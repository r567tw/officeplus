#!/usr/bin/env python3

import sys
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment


# =========================
# Configuration
# =========================

ADDED_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
DELETED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
MODIFIED_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")

HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")


# =========================
# Helpers
# =========================

def normalize_value(value):
    """
    Normalize Excel values before comparison.

    None is treated as an empty cell.
    """
    if value is None:
        return ""

    return value


def get_cell_value(ws, row, col):
    return normalize_value(ws.cell(row=row, column=col).value)


def compare_sheets(ws_old, ws_new):
    """
    Compare two worksheets cell by cell.

    Returns a list of differences:
    {
        "cell": "A1",
        "type": "ADDED" / "DELETED" / "MODIFIED",
        "old": old_value,
        "new": new_value
    }
    """

    max_row = max(ws_old.max_row, ws_new.max_row)
    max_col = max(ws_old.max_column, ws_new.max_column)

    differences = []

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            old_value = get_cell_value(ws_old, row, col)
            new_value = get_cell_value(ws_new, row, col)

            # Both empty -> no difference
            if old_value == "" and new_value == "":
                continue

            # Added
            if old_value == "" and new_value != "":
                differences.append({
                    "cell": ws_new.cell(row=row, column=col).coordinate,
                    "type": "ADDED",
                    "old": "",
                    "new": new_value,
                })

            # Deleted
            elif old_value != "" and new_value == "":
                differences.append({
                    "cell": ws_old.cell(row=row, column=col).coordinate,
                    "type": "DELETED",
                    "old": old_value,
                    "new": "",
                })

            # Modified
            elif old_value != new_value:
                differences.append({
                    "cell": ws_old.cell(row=row, column=col).coordinate,
                    "type": "MODIFIED",
                    "old": old_value,
                    "new": new_value,
                })

    return differences


def compare_workbooks(old_file, new_file):
    """
    Compare all worksheets in two Excel files.
    """

    wb_old = load_workbook(old_file, data_only=False)
    wb_new = load_workbook(new_file, data_only=False)

    differences = []

    old_sheets = set(wb_old.sheetnames)
    new_sheets = set(wb_new.sheetnames)

    # Deleted sheets
    for sheet_name in sorted(old_sheets - new_sheets):
        differences.append({
            "sheet": sheet_name,
            "cell": "",
            "type": "SHEET_DELETED",
            "old": "Sheet exists",
            "new": "",
        })

    # Added sheets
    for sheet_name in sorted(new_sheets - old_sheets):
        differences.append({
            "sheet": sheet_name,
            "cell": "",
            "type": "SHEET_ADDED",
            "old": "",
            "new": "Sheet exists",
        })

    # Compare common sheets
    for sheet_name in sorted(old_sheets & new_sheets):
        ws_old = wb_old[sheet_name]
        ws_new = wb_new[sheet_name]

        sheet_differences = compare_sheets(ws_old, ws_new)

        for diff in sheet_differences:
            differences.append({
                "sheet": sheet_name,
                **diff,
            })

    return differences


def create_diff_workbook(differences, output_file):
    """
    Create diff.xlsx with a summary sheet.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Diff"

    headers = [
        "Sheet",
        "Cell",
        "Type",
        "Old Value",
        "New Value",
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, diff in enumerate(differences, start=2):
        ws.cell(row=row_idx, column=1, value=diff["sheet"])
        ws.cell(row=row_idx, column=2, value=diff["cell"])
        ws.cell(row=row_idx, column=3, value=diff["type"])
        ws.cell(row=row_idx, column=4, value=diff["old"])
        ws.cell(row=row_idx, column=5, value=diff["new"])

        row_cells = ws[row_idx]

        if diff["type"] == "ADDED":
            fill = ADDED_FILL
        elif diff["type"] == "DELETED":
            fill = DELETED_FILL
        elif diff["type"] == "MODIFIED":
            fill = MODIFIED_FILL
        else:
            fill = None

        if fill:
            for cell in row_cells:
                cell.fill = fill

    # Freeze header
    ws.freeze_panes = "A2"

    # Enable filter
    ws.auto_filter.ref = ws.dimensions

    # Column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 40

    wb.save(output_file)


# =========================
# Main
# =========================

def main():
    if len(sys.argv) != 3:
        print(
            "Usage:\n"
            "  python excel_diff.py old.xlsx new.xlsx\n"
        )
        sys.exit(1)

    old_file = Path(sys.argv[1])
    new_file = Path(sys.argv[2])

    if not old_file.exists():
        print(f"Error: file not found: {old_file}")
        sys.exit(1)

    if not new_file.exists():
        print(f"Error: file not found: {new_file}")
        sys.exit(1)

    output_file = Path("diff.xlsx")

    print(f"Comparing:")
    print(f"  OLD: {old_file}")
    print(f"  NEW: {new_file}")

    differences = compare_workbooks(
        old_file,
        new_file,
    )

    create_diff_workbook(
        differences,
        output_file,
    )

    print()
    print(f"Found {len(differences)} difference(s).")
    print(f"Output: {output_file}")


if __name__ == "__main__":
    main()


